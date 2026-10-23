"""The generators the data refresh runs find their sources by rule and read
them by their own headers and wording (app/data_refresh.py). Offline: tiny
synthetic workbooks, questionnaires and listings."""

import importlib.util
import io
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl

from app.acceptance import held_out_beats_in_use

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"


def _script(name: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _book(sheets: dict[str, list[list]]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return openpyxl.load_workbook(io.BytesIO(buf.getvalue()), read_only=True, data_only=True)


def test_egrid_reads_the_us_total_output_co2e_rate_by_its_headers():
    egrid = _script("fetch_egrid")
    book = _book({
        "Contents": [["eGRID Summary Tables 2031"]],
        "Table 1": [
            ["1. Subregion Output Emission Rates (eGRID2031)"],
            ["acronym", "name", "Total output emission rates", "CH₄", "N₂O", "CO₂e", "Non-baseload", "CO₂e"],
            ["acronym", "name", "CO₂", "CH₄", "N₂O", "CO₂e", "CO₂", "CO₂e"],
            ["AKGD", "ASCC Alaska Grid", 899.6, 0.08, 0.01, 905.1, 1077.5, 1080.0],
            ["U.S.", None, 650.0, 0.05, 0.007, 652.4, 1100.0, 1105.0],
        ],
    })
    assert egrid.us_rate(book) == (652.4, 2031)
    page = '<a href="/system/files/documents/2032-06/summary_tables_rev1.xlsx">eGRID 2031 Summary Tables (xlsx)</a>'
    assert egrid.workbook_url(page) == "https://www.epa.gov/system/files/documents/2032-06/summary_tables_rev1.xlsx"


def test_maddison_release_is_the_newest_by_name_and_its_columns_are_read_by_header():
    gdp = _script("fetch_peer_gdp")
    items = [{"name": "Maddison Project Database 2023", "global_id": "doi:a"},
             {"name": "Maddison Project Database 2030", "global_id": "doi:b"},
             {"name": "Replication Data for Chapter 5", "global_id": "doi:c"}]
    assert gdp.newest_release(items) == (2030, "doi:b")
    book = _book({"Notes": [["About"], ["x"]],
                  "Full data": [["year", "countrycode", "country", "pop", "gdppc"],
                                [1950, "USA", "United States", 150, 15000.0], [1950, "FRA", "France", 40, 8000.0],
                                [1950, "XXX", "Elsewhere", 1, 100.0]]})
    assert list(gdp.per_capita_rows(book, {"USA", "FRA"})) == [("USA", 1950, 15000.0), ("FRA", 1950, 8000.0)]
    version = {"versionNumber": 2, "versionMinorNumber": 1, "files": [
        {"dataFile": {"id": 7, "filename": "mpd2030.xlsx"}}, {"dataFile": {"id": 8, "filename": "mpd2030.dta"}}],
        "metadataBlocks": {"citation": {"fields": [{"typeName": "publication", "value": [
            {"publicationCitation": {"value": "Bolt and van Zanden (2031)"}}]}]}}}
    assert gdp.release_file(version) == (7, "mpd2030.xlsx", "2.1", "Bolt and van Zanden (2031)")


def test_ces_surveys_are_found_by_name_and_their_items_mapped_by_wording():
    ces = _script("fetch_ces_approval")
    assert ces.survey_year("CES Common Content, 2025") == 2025
    assert ces.survey_year("2017 CCES Common Content") == 2017
    assert ces.survey_year("Cumulative CES Common Content") is None
    assert ces.survey_year("2020 - 2024 Cooperative Election Study Panel Survey") is None
    old = ces.wordings(["Do you approve of the way each is doing their job?", "CC24_312a", "President Biden",
                        "CC24_312f- Show if CurrentHouseName", "$CurrentHouseName",
                        "CC24_312g- Show if CurrentSen1Name", "$CurrentSen1Name",
                        "Generally speaking, do you think of yourself as a ...?", "CC24_400 party"])
    # A new survey renumbers and reorders the grid.
    new = ces.wordings(["Do you approve of the way each is doing their job?",
                        "CC26_318g- Show if $CurrentSen1Name", "CurrentSen1Name",
                        "CC26_318a President Smith", "CC26_318f- Show if $CurrentHouseName", "CurrentHouseName",
                        "Generally speaking, do you think of yourself as a ...?", "CC26_401 party"])

    def encode(texts):  # a stand-in model: letter counts, unit length
        vecs = np.array([[t.count(c) for c in "abcdefghijklmnopqrstuvwxyz0123456789$"] for t in texts], float)
        return vecs / np.linalg.norm(vecs, axis=1, keepdims=True)

    assert ces.map_items(old, new, ["CC24_312f", "CC24_312g"], encode) == {"CC24_312f": "CC26_318f",
                                                                          "CC24_312g": "CC26_318g"}
    assert "not in the questionnaire" in ces.map_items(old, new, ["CC24_999"], encode)
    assert ces.district_column(["cdid118", "cdid119", "cdid120"], datetime(2025, 11, 1)) == "cdid119"
    first, last = ces.fielded([{"starttime": "2.079984094E12"}, {"starttime": "2025-10-02T10:00:00"}])
    assert (first.year, first.month, last.month) == (2025, 10, 11)
    rows = [{"Q": "1"}, {"Q": "3"}, {"Q": "5"}, {"Q": ""}]
    assert ces.plausible(rows, [("Q",)]) is None
    assert "approve" in ces.plausible([{"Q": "1"}] * 20 + [{"Q": "3"}], [("Q",)])


def test_the_forecast_priors_are_kept_unless_they_win_on_the_forward_check():
    fit = _script("fit_forecast_model")
    win = fit.priors_comparison({"senate": ([-3.0, -3.1, -2.9, -3.2], [-2.0, -2.1, -2.2, -1.9]),
                                 "house": ([-3.0] * 5, [-2.5, -2.6, -2.4, -2.5, -2.6])})
    assert held_out_beats_in_use({"_comparison": win}, {}) is None
    noise = fit.priors_comparison({"senate": ([-3.0, -3.1, -2.9], [-3.1, -3.0, -2.95]),
                                   "house": ([-3.0, -3.0], [-3.0, -3.0])})
    assert "within noise" in held_out_beats_in_use({"_comparison": noise}, {})
